"""Generador de la guia de desplegament per l'equip de Sistemes.

Genera `docs/Desplegament_SAP_Sistemes.docx` (+ `.pdf`): donar un nom propi a la
variant SAP (`comandessap.agrienergia.local`) sense tocar la URL que ja fan servir
els usuaris de la versio Kais.

Public objectiu: administradors de sistemes, no desenvolupadors. Cada comandament
va acompanyat de que fa i de com verificar que ha anat be.

Reutilitza els helpers i la paleta de `build_proposta_sap.py`.

Us:
    python scripts/build_guia_sistemes.py
"""
from __future__ import annotations

import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_proposta_sap import (  # noqa: E402
    COLOR_MUTED,
    COLOR_SUBTITOL,
    COLOR_TITOL,
    add_bullet,
    add_code_block,
    add_heading,
    add_info_box,
    add_page_break,
    add_para,
    add_taula,
)

COLOR_AVIS_BG = "FFF4CE"
COLOR_OK_BG = "E8F5E9"

IP = "192.168.11.244"
SERVIDOR = "ae01farwebsrv"
APP_DIR = "/var/www/comandes-venda-sap"
URL_SAP = "comandessap.agrienergia.local"
URL_KAIS = "comandes.agrienergia.local"
CONF_SAP = "comandes-venda-sap.conf"
CONF_KAIS = "comandes-venda.conf"


def add_avis(doc, text):
    add_info_box(doc, text, bg=COLOR_AVIS_BG, icon="⚠")


def add_ok(doc, text):
    add_info_box(doc, text, bg=COLOR_OK_BG, icon="✔")


def add_pas(doc, numero, titol):
    add_heading(doc, f"Pas {numero} — {titol}", level=2)


def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ================= PORTADA =================
    for _ in range(7):
        doc.add_paragraph()

    add_para(doc, "GUIA DE DESPLEGAMENT", bold=True, size=24, color=COLOR_TITOL,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Motor d'Embalatges — versió SAP Business One", bold=True,
             size=17, color=COLOR_SUBTITOL, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=20)
    add_para(doc, f"Donar-li un nom propi: http://{URL_SAP}/",
             italic=True, size=12.5, color=COLOR_MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)

    for _ in range(5):
        doc.add_paragraph()

    add_para(doc, "Destinatari: equip de Sistemes", bold=True, size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, f"Data: {date.today().strftime('%d/%m/%Y')}", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Empresa: Agrienergia — Farinera Coromina", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Autor: Oscar Hijazo", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    add_page_break(doc)

    # ================= 1. RESUM =================
    add_heading(doc, "1. Resum executiu", level=1)

    add_para(doc,
             f"Al servidor {SERVIDOR} conviuen dues versions de l'aplicació de "
             "càlcul d'embalatges de comandes de venda. La que està en producció a "
             f"http://{URL_KAIS}/ llegeix les dades del sistema antic (Kais). La "
             "nova llegeix de SAP Business One i ja corre al mateix servidor.")

    add_para(doc,
             "La versió nova només és accessible ara mateix escrivint la IP i el "
             f"port (http://{IP}:5002/), cosa incòmoda per als usuaris i fràgil si "
             "algun dia canvia la IP. Volem donar-li un nom propi.")

    add_ok(doc, f"La versió en producció NO es toca: {URL_KAIS} es queda amb el "
                "mateix nom, la mateixa configuració i el mateix servei. Els seus "
                "usuaris no noten res.")

    add_heading(doc, "Què cal fer", level=2)
    add_taula(doc,
              ["Tasca", "Responsable", "Temps", "Estat"],
              [
                  ["Activar el VirtualHost d'Apache", "Oscar", "5 min",
                   "Es pot fer ja"],
                  [f"Registre DNS {URL_SAP} (apartat 3)", "Sistemes", "5 min",
                   "Pendent"],
                  ["Verificació conjunta (apartat 5)", "Tots dos", "5 min",
                   "Pendent del DNS"],
                  ["Comunicar la URL als usuaris", "Oscar", "—", "Pendent"],
              ],
              amples=[6.5, 3.0, 2.0, 4.5])

    add_info_box(doc,
                 "L'ordre no importa: el VirtualHost es pot activar abans que "
                 "existeixi el registre DNS, perquè Apache només encamina pel nom "
                 "de host que li arriba i cap client enviarà aquest nom fins que "
                 "resolgui. Així es pot deixar tot muntat i verificat esperant "
                 "només l'alta.")

    add_page_break(doc)

    # ================= 2. ARQUITECTURA =================
    add_heading(doc, "2. Com està muntat el servidor", level=1)

    add_para(doc,
             "Apache fa de proxy invers al port 80 amb VirtualHosts per nom. Cada "
             "aplicació és un procés Gunicorn escoltant a 127.0.0.1 al seu port, i "
             "Apache l'encamina segons el nom de host de la petició. Al servidor hi "
             "conviuen cinc aplicacions.")

    add_code_block(doc,
                   "                 Apache  *:80  (NameVirtualHost)\n"
                   "                 ---------------------------------\n"
                   "  agrupacions.agrienergia.local --> agrupacio-carregues.conf   (default)\n"
                   f"  {URL_KAIS} -----> {CONF_KAIS}      --> 127.0.0.1:5001  Kais\n"
                   "  fitxesfc.agrienergia.local ------> fitxes-tecniques.conf\n"
                   "  labfc.agrienergia.local ---------> labfc.conf\n"
                   "  visitesfc.agrienergia.local -----> visites.conf\n"
                   "\n"
                   f"  {URL_SAP} --> {CONF_SAP}  --> 127.0.0.1:5002  SAP\n"
                   "                                             (per activar)")

    add_para(doc,
             "L'única línia nova és l'última. Les altres cinc no es toquen.",
             size=10, color=COLOR_MUTED)

    add_heading(doc, "Estat de la versió SAP", level=2)
    add_taula(doc,
              ["Element", "Valor"],
              [
                  ["Servei systemd", "comandes-venda-sap.service"],
                  ["Servidor web", "Gunicorn (gthread, 2 processos × 4 fils)"],
                  ["Socket", "0.0.0.0:5002"],
                  ["Directori", APP_DIR],
                  ["VirtualHost d'Apache", f"{CONF_SAP} (preparat, per activar)"],
              ],
              amples=[5.0, 11.0])

    add_para(doc,
             "El socket és 0.0.0.0 i no 127.0.0.1 perquè el botó «Calcular "
             f"embalatges» de SAP Business One crida http://{IP}:5002/ per IP "
             "directa. Aquesta via ha de seguir funcionant; veure l'apartat 7.")

    add_page_break(doc)

    # ================= 3. PETICIÓ =================
    add_heading(doc, "3. Què demanem a Sistemes", level=1)

    add_para(doc, "Un sol registre DNS:")

    add_taula(doc,
              ["Nom", "Tipus", "Valor", "TTL"],
              [[URL_SAP, "A", IP, "300"]],
              amples=[7.5, 2.0, 4.0, 2.5])

    add_info_box(doc,
                 f"És la MATEIXA IP que ja fa servir {URL_KAIS}. El que separa les "
                 "dues aplicacions és el nom de host que llegeix Apache, no la "
                 "xarxa. No cal cap IP nova.")

    add_heading(doc, "Què NO cal fer", level=2)
    add_bullet(doc, f"No cal tocar el registre {URL_KAIS}.")
    add_bullet(doc, "No cal cap IP nova ni cap canvi a la configuració de xarxa.")
    add_bullet(doc, "No cal obrir ports al tallafocs: és el port 80 que Apache ja "
                    "fa servir.")
    add_bullet(doc, "No cal aturar ni reconfigurar comandes-venda.service, "
                    "l'aplicació en producció.")

    add_para(doc, "Verificació un cop donat d'alta:")
    add_code_block(doc, f"nslookup {URL_SAP}")

    add_page_break(doc)

    # ================= 4. ACTIVACIÓ =================
    add_heading(doc, "4. Activar el VirtualHost", level=1)

    add_para(doc,
             "Ho executa l'Oscar per SSH. Es documenta perquè Sistemes ho pugui "
             "seguir o repetir. No cal esperar el DNS.")

    add_pas(doc, 1, "Portar el fitxer al servidor")
    add_code_block(doc, f"sudo bash {APP_DIR}/deploy.sh")

    add_pas(doc, 2, "Activar-lo")
    add_code_block(doc,
                   f"sudo cp {APP_DIR}/deploy/apache/{CONF_SAP} \\\n"
                   "    /etc/apache2/sites-available/\n"
                   f"sudo a2ensite {CONF_SAP}\n"
                   "sudo apachectl configtest\n"
                   "sudo systemctl reload apache2")

    add_avis(doc, "Si «apachectl configtest» no diu «Syntax OK», NO recarregar. Un "
                  "«reload» amb configuració invàlida deixa Apache servint la "
                  "configuració antiga, però un «restart» posterior el tombaria i "
                  "s'emportaria les altres cinc aplicacions del servidor.")

    add_para(doc,
             "Es fa «reload» i no «restart» precisament per no tallar connexions "
             "actives de cap aplicació.", size=10, color=COLOR_MUTED)

    add_heading(doc, "Contingut del VirtualHost", level=2)
    add_para(doc,
             "Calca el patró del de Kais: mateix «AllowEncodedSlashes NoDecode» "
             "(cal per als números de sèrie amb caràcters especials), mateix "
             "«ProxyPreserveHost», i els fitxers estàtics servits per Apache en "
             "lloc de passar per Gunicorn. L'única diferència és el port del "
             "backend (5002 en lloc de 5001) i el directori.")

    # ================= 5. VERIFICACIÓ =================
    add_heading(doc, "5. Verificació", level=1)

    add_heading(doc, "5.1 Al servidor, abans i tot que hi hagi DNS", level=2)
    add_para(doc,
             "Amb la capçalera «Host» es pot provar tot el camí sense esperar que "
             "el nom resolgui:")
    add_code_block(doc,
                   f'curl -sS -H "Host: {URL_SAP}" http://127.0.0.1/ajuda | head -5\n'
                   "\n"
                   'curl -sS -o /dev/null -w "%{http_code} %{content_type}\\n" \\\n'
                   f'    -H "Host: {URL_SAP}" http://127.0.0.1/static/css/style.css\n'
                   "\n"
                   f'curl -sS -H "Host: {URL_KAIS}" http://127.0.0.1/ | head -5\n'
                   "\n"
                   "sudo apachectl -S | grep comandes")

    add_taula(doc,
              ["Comprovació", "Resultat esperat"],
              [
                  ["Pàgina d'ajuda de SAP", "HTML amb «(SAP)» al títol"],
                  ["Fitxer estàtic", "200 text/css — el serveix Apache"],
                  ["Kais", "HTML de Kais, com sempre"],
                  ["apachectl -S", "els dos VirtualHosts registrats"],
              ],
              amples=[6.5, 9.5])

    add_heading(doc, "5.2 Des d'un PC Windows, un cop hi hagi DNS", level=2)
    add_code_block(doc,
                   f"nslookup {URL_SAP}\n"
                   f"curl.exe http://{URL_SAP}/ajuda\n"
                   f"curl.exe http://{URL_KAIS}/\n"
                   f"curl.exe http://{IP}:5002/api/admin/versio")

    add_ok(doc, "Criteri d'acceptació: les tres URLs responen. L'última és la que "
                "garanteix que el botó «Calcular embalatges» de SAP Business One "
                "segueix funcionant.")

    add_heading(doc, "5.3 Prova funcional a SAP Business One", level=2)
    add_bullet(doc, "Obrir el SAP Fat Client i entrar a Comanda de venda "
                    "(una comanda esborrany).")
    add_bullet(doc, "Clicar el botó «Calcular embalatges».")
    add_bullet(doc, "Comprovar el missatge d'èxit a la barra d'estat.")
    add_bullet(doc, "Comprovar que apareixen les línies de palet a la comanda.")

    add_page_break(doc)

    # ================= 6. MARXA ENRERE =================
    add_heading(doc, "6. Marxa enrere i diagnòstic", level=1)

    add_para(doc,
             "La versió en producció no es toca en cap moment, per tant no té marxa "
             "enrere ni la necessita. Tot el que es pot desfer afecta només "
             "l'aplicació SAP.")

    add_code_block(doc,
                   f"sudo a2dissite {CONF_SAP}\n"
                   "sudo apachectl configtest && sudo systemctl reload apache2")
    add_para(doc,
             f"{URL_SAP} deixa de respondre. L'aplicació segueix accessible per "
             f"{IP}:5002 i el botó de SAP B1 no se n'assabenta.",
             size=10, color=COLOR_MUTED)

    add_heading(doc, "Logs", level=2)
    add_code_block(doc,
                   "sudo journalctl -u comandes-venda-sap -n 50\n"
                   f"sudo tail -50 {APP_DIR}/error.log\n"
                   "sudo tail -50 /var/log/apache2/comandes-venda-sap-error.log")

    add_taula(doc,
              ["Símptoma", "Causa", "Solució"],
              [
                  ["502 / 503 al navegador",
                   "Gunicorn aturat o no escolta al 5002",
                   "systemctl status comandes-venda-sap"],
                  ["El nom serveix una altra aplicació",
                   "VirtualHost no activat o Apache no recarregat",
                   "sudo apachectl -S | grep comandes"],
                  ["«Syntax error» al configtest",
                   "Error al fitxer del VirtualHost",
                   "No recarregar; revisar el fitxer"],
                  ["El nom no resol",
                   "Registre DNS no donat d'alta o no propagat",
                   f"nslookup {URL_SAP}"],
                  ["Els estàtics donen 404",
                   "Ruta de l'Alias incorrecta",
                   f"ls {APP_DIR}/static/"],
              ],
              amples=[4.5, 5.5, 6.0])

    add_page_break(doc)

    # ================= 7. BOTÓ B1UP =================
    add_heading(doc, "7. El botó de SAP Business One", level=1)

    add_para(doc,
             "Dins de SAP Business One, al formulari Comanda de venda, hi ha un "
             "botó «Calcular embalatges» (configurat amb B1UP) que crida "
             f"l'aplicació per HTTP a http://{IP}:5002/api/afegir-palets/.")

    add_ok(doc, "El canvi d'aquesta guia no l'afecta: crida l'aplicació per IP i "
                "port directes, no per nom. No cal coordinar el consultor de B1UP.")

    add_para(doc,
             "Un cop el nom nou funcioni, el consultor el pot reapuntar a "
             f"http://{URL_SAP}/api/afegir-palets/ per treure la IP del codi de "
             "B1UP. Llavors es pot tancar Gunicorn darrere d'Apache, que és la "
             "configuració desitjable:")
    add_code_block(doc,
                   f"sudo BIND_ADDR=127.0.0.1:5002 bash {APP_DIR}/deploy.sh --reinstall-service")
    add_para(doc, "Res d'això és urgent ni bloqueja el nom nou.",
             size=10, color=COLOR_MUTED)

    # ================= 8. MANTENIMENT =================
    add_heading(doc, "8. Manteniment", level=1)

    add_heading(doc, "Actualitzar l'aplicació", level=2)
    add_code_block(doc,
                   "# Codi al dia (git pull + reinici del servei)\n"
                   f"sudo bash {APP_DIR}/deploy.sh\n"
                   "\n"
                   "# Reescriure la definició del servei systemd. Repetible.\n"
                   f"sudo bash {APP_DIR}/deploy.sh --reinstall-service")

    add_para(doc,
             "El segon comandament existeix perquè el primer no toca la definició "
             "del servei. Per això l'aplicació havia quedat mesos amb el servidor "
             "de desenvolupament de Flask tot i que el repositori ja definia "
             "Gunicorn: cap actualització hi arribava.", size=10, color=COLOR_MUTED)

    add_heading(doc, "Comandaments habituals", level=2)
    add_code_block(doc,
                   "sudo systemctl status comandes-venda-sap\n"
                   "sudo systemctl restart comandes-venda-sap\n"
                   "sudo journalctl -u comandes-venda-sap -f")

    add_heading(doc, "Logs i rotació", level=2)
    add_para(doc,
             f"Gunicorn escriu a {APP_DIR}/access.log i error.log, amb rotació "
             "setmanal i 4 setmanes de retenció "
             "(/etc/logrotate.d/comandes-venda-sap). Apache, a "
             "/var/log/apache2/comandes-venda-sap-*.log. Comprovació el dilluns: "
             "els fitxers de la setmana anterior han de tenir extensió .1.gz.")

    add_heading(doc, "Les dues aplicacions d'embalatges", level=2)
    add_taula(doc,
              ["Versió", "URL", "Servei systemd", "Port intern"],
              [
                  ["Kais", URL_KAIS, "comandes-venda.service", "5001"],
                  ["SAP", URL_SAP, "comandes-venda-sap.service", "5002"],
              ],
              amples=[2.0, 6.5, 5.0, 2.5])

    add_avis(doc, "El servei comandes-venda.service (versió Kais) ha de quedar "
                  "operatiu fins al novembre de 2026. Cap pas d'aquesta guia el "
                  "toca ni l'atura.")

    # ================= CONTACTE =================
    add_heading(doc, "Contacte", level=1)
    add_para(doc, "Oscar Hijazo — ohijazo@agrienergia.com", bold=True)
    add_para(doc,
             "Documentació tècnica al repositori de l'aplicació: deploy/README.md "
             "(arquitectura de xarxa del servidor), docs/runbook_url_propia_sap.md "
             "(procediment detallat), docs/peticio_dns_sistemes.md (la petició de "
             "l'apartat 3 en text pla).", size=10, color=COLOR_MUTED)

    return doc


SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def _exporta_pdf(docx_path, pdf_path):
    """Converteix a PDF. LibreOffice primer (no depen que Word estigui tancat)."""
    if os.path.exists(SOFFICE):
        import shutil
        import subprocess
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            res = subprocess.run(
                [SOFFICE, "--headless", "--norestore", "--convert-to", "pdf",
                 "--outdir", tmp, docx_path],
                capture_output=True, text=True)
            generat = os.path.join(tmp, os.path.basename(docx_path)[:-5] + ".pdf")
            if res.returncode == 0 and os.path.exists(generat):
                shutil.copy2(generat, pdf_path)
                return True
            print("AVIS LibreOffice:", res.stderr.strip() or res.stdout.strip())

    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        return True
    except Exception as e:  # noqa: BLE001 - el .docx ja esta desat; el PDF es opcional
        print("AVIS docx2pdf: {} (tanca el Word i torna-ho a provar)".format(e))
        return False


def main():
    docs_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
    os.makedirs(docs_dir, exist_ok=True)
    docx_path = os.path.abspath(os.path.join(docs_dir, "Desplegament_SAP_Sistemes.docx"))
    pdf_path = docx_path[:-5] + ".pdf"

    doc = build_document()
    doc.save(docx_path)
    print("Generat: {} ({:.1f} KB)".format(docx_path, os.path.getsize(docx_path) / 1024))

    if _exporta_pdf(docx_path, pdf_path):
        print("Generat: {} ({:.1f} KB)".format(pdf_path, os.path.getsize(pdf_path) / 1024))
    else:
        print("AVIS: no s'ha pogut generar el PDF. Obre el .docx i desa'l com a PDF.")


if __name__ == "__main__":
    main()
